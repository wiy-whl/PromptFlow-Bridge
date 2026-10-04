# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from pathlib import Path


BLUE = RGBColor(0x2F, 0x6F, 0xED)
BLACK = RGBColor(0x11, 0x11, 0x11)
GRAY = RGBColor(0x44, 0x44, 0x44)
LINE = "CFD6E0"


def set_run_font(run, size=10.5, bold=False, color=BLACK, name="微软雅黑"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def set_paragraph_format(p, space_before=0, space_after=4, line=1.35, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align


def add_bottom_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), LINE)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_section_title(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, space_before=10, space_after=6, line=1.2)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, color=BLACK)
    add_bottom_border(p)
    return p


def add_mixed_paragraph(doc, parts, space_before=0, space_after=3, justify=True):
    """parts: list of (text, bold, color|None, size)"""
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        space_before=space_before,
        space_after=space_after,
        line=1.4,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY if justify else None,
    )
    for text, bold, color, size in parts:
        run = p.add_run(text)
        set_run_font(run, size=size or 10.5, bold=bold, color=color or BLACK)
    return p


def add_left_right_row(doc, left, right, left_bold=False, right_bold=True, size=10.5):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    table.allow_autofit = True
    cell_l, cell_r = table.rows[0].cells
    # clear default paras
    for cell, text, bold, align in (
        (cell_l, left, left_bold, WD_ALIGN_PARAGRAPH.LEFT),
        (cell_r, right, right_bold, WD_ALIGN_PARAGRAPH.RIGHT),
    ):
        cell.text = ""
        p = cell.paragraphs[0]
        set_paragraph_format(p, space_before=0, space_after=0, line=1.25, align=align)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=BLACK)
        # remove cell margins a bit
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcMar = OxmlElement("w:tcMar")
        for m in ("top", "left", "bottom", "right"):
            node = OxmlElement(f"w:{m}")
            node.set(qn("w:w"), "0")
            node.set(qn("w:type"), "dxa")
            tcMar.append(node)
        tcPr.append(tcMar)
    # hide borders
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        borders.append(elem)
    tblPr.append(borders)
    return table


def build():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    # ===== Header: name + basic info + photo placeholder =====
    head = doc.add_table(rows=1, cols=2)
    left, right = head.rows[0].cells
    # hide borders
    tbl = head._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        borders.append(elem)
    tblPr.append(borders)

    left.width = Cm(14.5)
    right.width = Cm(2.8)

    left.text = ""
    p = left.paragraphs[0]
    set_paragraph_format(p, space_before=0, space_after=8, line=1.1)
    run = p.add_run("傅怡")
    set_run_font(run, size=26, bold=True, color=BLACK)

    info = [
        ("求职：UI / 视觉设计师（C端）", "城市：深圳"),
        ("电话：13177850203", "邮箱：3528544024@qq.com"),
        ("届别：2027届（可实习）", "作品集：制作中，可索取"),
    ]
    info_table = left.add_table(rows=3, cols=2)
    itbl = info_table._tbl
    itblPr = itbl.tblPr if itbl.tblPr is not None else OxmlElement("w:tblPr")
    iborders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        iborders.append(elem)
    itblPr.append(iborders)
    for i, (a, b) in enumerate(info):
        for j, text in enumerate((a, b)):
            cell = info_table.rows[i].cells[j]
            cell.text = ""
            cp = cell.paragraphs[0]
            set_paragraph_format(cp, space_before=0, space_after=2, line=1.25)
            set_run_font(cp.add_run(text), size=10.5, bold=False, color=BLACK)

    right.text = ""
    rp = right.paragraphs[0]
    set_paragraph_format(rp, space_before=0, space_after=0, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    # photo box via border paragraph
    box = right.add_paragraph()
    set_paragraph_format(box, space_before=0, space_after=0, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    # simulate placeholder text
    set_run_font(box.add_run("一寸照\n可替换"), size=9, bold=False, color=RGBColor(0x98, 0xA2, 0xB3))
    # add border to right cell
    tc = right._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:color"), "D0D5DD")
        tcBorders.append(elem)
    tcPr.append(tcBorders)

    # ===== 教育经历 =====
    add_section_title(doc, "教育经历")
    add_left_right_row(doc, "2023.09 – 2027.07", "豫章师范学院")
    add_mixed_paragraph(
        doc,
        [("视觉传达设计｜本科", False, GRAY, 10.5)],
        space_after=2,
        justify=False,
    )

    # ===== 专业技能 =====
    add_section_title(doc, "专业技能")
    skills = [
        [
            ("【界面设计】", True, BLACK, 10.5),
            ("熟练使用 ", False, BLACK, 10.5),
            ("Figma", True, BLUE, 10.5),
            (" 完成 C 端界面设计，掌握 Auto Layout、组件化与基础设计规范（色彩 / 字体 / 间距），能输出便于标注与还原的界面稿。", False, BLACK, 10.5),
        ],
        [
            ("【视觉执行】", True, BLACK, 10.5),
            ("熟练使用 ", False, BLACK, 10.5),
            ("Photoshop、Illustrator", True, BLUE, 10.5),
            ("，具备主 KV、图标、运营物料与样机呈现能力；具备品牌 IP / 角色视觉延展经验。", False, BLACK, 10.5),
        ],
        [
            ("【AI 设计工具】", True, BLACK, 10.5),
            ("熟练使用 ", False, BLACK, 10.5),
            ("ChatGPT、Gemini、即梦、哩布哩布", True, BLUE, 10.5),
            (" 等主流 AI：用大模型辅助信息架构与文案梳理，用图像生成做风格探索与素材发散，再回到设计工具完成筛选、规范统一与界面落地。", False, BLACK, 10.5),
        ],
        [
            ("【产品协作意识】", True, BLACK, 10.5),
            ("能梳理首页到核心功能的信息层级，关注状态表达与多屏一致性，适配互联网团队协作节奏。", False, BLACK, 10.5),
        ],
    ]
    for parts in skills:
        add_mixed_paragraph(doc, parts, space_after=4)

    # ===== 项目经历 =====
    add_section_title(doc, "项目经历")

    # Project 1
    add_left_right_row(doc, "2025", "大理·文化信息体验小程序 / Web")
    add_mixed_paragraph(doc, [("界面与视觉设计（课程项目 / 独立完成）", False, GRAY, 10)], justify=False, space_after=2)
    add_mixed_paragraph(
        doc,
        [
            ("项目概述：", True, BLACK, 10.5),
            ("面向年轻游客的大理文旅数字化体验，独立完成微信小程序核心界面与 Web 专题页，从文化符号提取到多端界面落地形成完整闭环。", False, BLACK, 10.5),
        ],
        space_after=2,
    )
    add_mixed_paragraph(
        doc,
        [
            ("【信息架构与主路径】", True, BLACK, 10.5),
            ("梳理文化 / 景点 / 文创 / 攻略等模块，完成引导进入到内容详情的主路径设计，降低信息查找成本。", False, BLACK, 10.5),
        ],
        space_after=2,
    )
    add_mixed_paragraph(
        doc,
        [
            ("【视觉体系与规范】", True, BLACK, 10.5),
            ("提取扎染、瓦猫等符号落地 Logo、图标与蓝白视觉语言；结合 ", False, BLACK, 10.5),
            ("ChatGPT / 即梦 / 哩布哩布", True, BLUE, 10.5),
            (" 辅助结构梳理与风格探索，并在设计工具中完成规范统一与落地。", False, BLACK, 10.5),
        ],
        space_after=8,
    )

    # Project 2
    add_left_right_row(doc, "2025", "娃哈哈 AD 钙 · 品牌年轻化数字体验")
    add_mixed_paragraph(doc, [("品牌视觉 / 界面视觉（个人练习，非官方项目）", False, GRAY, 10)], justify=False, space_after=2)
    add_mixed_paragraph(
        doc,
        [
            ("项目概述：", True, BLACK, 10.5),
            ("围绕儿童营养与亲子场景，完成品牌年轻化方向的视觉探索，输出品牌调性、图标体系与 App 核心界面视觉，强化 IP 与情感化表达。", False, BLACK, 10.5),
        ],
        space_after=2,
    )
    add_mixed_paragraph(
        doc,
        [
            ("【品牌与 IP 表达】", True, BLACK, 10.5),
            ("建立色板与 ", False, BLACK, 10.5),
            ("3D 图标", True, BLUE, 10.5),
            (" 体系，产出主视觉 KV 与功能展示页，突出年轻化、亲和的品牌气质。", False, BLACK, 10.5),
        ],
        space_after=2,
    )
    add_mixed_paragraph(
        doc,
        [
            ("【界面转化】", True, BLACK, 10.5),
            ("结合竞品与信息层级整理，将产品展示、营养教育、会员互动等模块转化为可落地的首页及核心功能布局。", False, BLACK, 10.5),
        ],
        space_after=8,
    )

    # Project 3
    add_left_right_row(doc, "2025", "查件 · 快递查询工具 App UI")
    add_mixed_paragraph(doc, [("UI 设计（个人项目 / 独立完成）", False, GRAY, 10)], justify=False, space_after=2)
    add_mixed_paragraph(
        doc,
        [
            ("项目概述：", True, BLACK, 10.5),
            ("面向高频查件场景，独立完成工具类 App 核心界面与精简设计规范，强调清晰、克制、可开发还原的产品视觉。", False, BLACK, 10.5),
        ],
        space_after=2,
    )
    add_mixed_paragraph(
        doc,
        [
            ("【核心界面】", True, BLACK, 10.5),
            ("设计包裹列表、物流详情时间轴、添加快递、个人中心等页面，针对在途 / 派送 / 签收等状态区分信息层级。", False, BLACK, 10.5),
        ],
        space_after=2,
    )
    add_mixed_paragraph(
        doc,
        [
            ("【组件一致性】", True, BLACK, 10.5),
            ("建立列表卡片、状态标签、时间轴等组件，统一多屏视觉语言，提升界面完整度与落地效率。", False, BLACK, 10.5),
        ],
        space_after=6,
    )

    # ===== 自我评价 =====
    add_section_title(doc, "自我评价")
    add_mixed_paragraph(
        doc,
        [
            ("视觉传达专业在读，具备排版、色彩与品牌视觉基础，并持续向 ", False, BLACK, 10.5),
            ("C 端 UI / 产品视觉", True, BLUE, 10.5),
            (" 转化；能将 ", False, BLACK, 10.5),
            ("ChatGPT、Gemini、即梦、哩布哩布", True, BLUE, 10.5),
            (" 融入设计工作流，重视规范与信息层级，能以完整项目呈现作品。求职方向为互联网 UI / 视觉设计（含 2027 届实习），可快速融入团队协作。", False, BLACK, 10.5),
        ],
        space_after=0,
    )

    out_dir = Path(r"d:\AI\chajian\Promptbridge\resume")
    out1 = out_dir / "傅怡-UI设计师-应届实习.docx"
    out2 = out_dir / "Fuyi-UI-Designer-Intern.docx"
    doc.save(out1)
    doc.save(out2)
    print("saved", out1)
    print("saved", out2)


if __name__ == "__main__":
    build()
